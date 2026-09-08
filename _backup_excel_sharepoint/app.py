# -*- coding: utf-8 -*-
"""
NSIA Assurances Togo - Call Center Dashboard
Application Streamlit pour la saisie, le suivi et l'analyse des appels.
Fichier Excel central: BDD_CALL_CENTER.xlsm
Colonnes calées sur la base NSIA existante.
"""

import streamlit as st
import pandas as pd
import plotly.express as px
from openpyxl import load_workbook
from datetime import datetime, date
import os

from points_de_vente import POINTS_DE_VENTE as DEFAULT_POINTS_DE_VENTE

# ============================================================================
# CONFIGURATION & THEME
# ============================================================================

NSIA = {
    "navy": "#202C54",
    "gold": "#CF9A06",
    "white": "#FFFFFF",
    "silver": "#ACA38B",
    "dark_gray": "#605C57",
    "blue_black": "#0D1222",
    "light_bg": "#F5F5F7",
    "card_bg": "#FFFFFF",
    "border": "#E0E0E3",
    "text_primary": "#202C54",
    "text_secondary": "#605C57",
    "success": "#2E7D32",
    "error": "#C62828",
    "warning": "#F9A825",
}

st.set_page_config(
    page_title="NSIA Call Center",
    page_icon="📞",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    f"""
<style>
/* Base */
.stApp {{
    background-color: {NSIA['blue_black']};
    color: {NSIA['silver']};
    font-family: 'Segoe UI', 'Inter', Roboto, sans-serif;
}}

/* Sidebar */
section[data-testid="stSidebar"] {{
    background: linear-gradient(180deg, {NSIA['navy']} 0%, {NSIA['blue_black']} 100%);
    color: {NSIA['white']};
}}
section[data-testid="stSidebar"] .stRadio > label,
section[data-testid="stSidebar"] .stCaption,
section[data-testid="stSidebar"] .stMarkdown {{
    color: {NSIA['white']} !important;
}}
section[data-testid="stSidebar"] h1,
section[data-testid="stSidebar"] h2,
section[data-testid="stSidebar"] h3 {{
    color: {NSIA['gold']} !important;
}}
section[data-testid="stSidebar"] hr {{
    border-color: rgba(207,154,6,0.25);
}}

/* Typography */
h1, h2, h3, h4, h5 {{
    color: {NSIA['white']};
    font-weight: 600;
    letter-spacing: -0.005em;
}}
h1 {{ font-size: 1.75rem; }}
h2 {{ font-size: 1.35rem; }}
h3 {{ font-size: 1.1rem; }}

/* Cartes / formulaires */
div[data-testid="stVerticalBlock"] > div:has(> div[data-testid="stForm"]) {{
    background-color: rgba(13,18,34,0.75);
    padding: 1.1rem 1.25rem;
    border-radius: 0.9rem;
    border: 1px solid rgba(172,163,139,0.15);
    box-shadow: 0 1px 6px rgba(0,0,0,0.35);
    margin-top: 0.6rem;
}}
div[data-testid="stHorizontalBlock"] > div {{
    gap: 0.65rem;
}}

/* Métriques */
div[data-testid="stMetric"] {{
    background-color: rgba(32,44,84,0.35);
    padding: 0.9rem 1.1rem;
    border-radius: 0.8rem;
    border-left: 3px solid {NSIA['gold']};
    box-shadow: 0 1px 4px rgba(0,0,0,0.25);
}}
div[data-testid="stMetric"] label {{
    color: {NSIA['silver']};
    font-size: 0.78rem;
    text-transform: uppercase;
    letter-spacing: 0.04em;
}}
div[data-testid="stMetric"] div[data-testid="stMetricValue"] {{
    color: {NSIA['white']};
    font-weight: 700;
    font-size: 1.25rem;
}}

/* Boutons */
button[kind="primary"] {{
    background-color: {NSIA['gold']};
    color: {NSIA['navy']};
    border: none;
    border-radius: 0.5rem;
    padding: 0.65rem 0.9rem;
    font-weight: 600;
    letter-spacing: 0.01em;
    box-shadow: 0 1px 5px rgba(207,154,6,0.25);
}}
button[kind="primary"]:hover {{
    background-color: #b8850a;
}}
button[kind="secondary"] {{
    background-color: rgba(255,255,255,0.08);
    color: {NSIA['white']};
    border: 1px solid rgba(172,163,139,0.35);
    border-radius: 0.5rem;
    padding: 0.65rem 0.9rem;
    font-weight: 600;
}}
button[kind="secondary"]:hover {{
    background-color: rgba(255,255,255,0.13);
    border-color: {NSIA['silver']};
}}

/* Chips campagne */
.chip {{
    display: inline-flex;
    align-items: center;
    justify-content: center;
    padding: 0.55rem 0.9rem;
    border-radius: 999px;
    border: 1px solid {NSIA['gold']};
    background: {NSIA['gold']};
    color: {NSIA['dark_gray']};
    font-weight: 700;
    transition: all 0.2s ease;
    margin: 0.2rem 0.3rem 0.2rem 0;
    font-size: 0.92rem;
    cursor: pointer;
    user-select: none;
}}
.chip:hover {{
    background-color: #b8850a;
    border-color: #b8850a;
    color: {NSIA['navy']};
}}
.chip.active {{
    background-color: {NSIA['navy']};
    color: {NSIA['white']};
    border-color: {NSIA['gold']};
    box-shadow: 0 1px 6px rgba(32,44,84,0.35);
}}

/* Dataframe */
div[data-testid="stDataFrame"] {{
    border-radius: 0.8rem;
    overflow: hidden;
    border: 1px solid rgba(172,163,139,0.25);
    background-color: rgba(13,18,34,0.6);
}}

/* Inputs */
div[data-testid="stSelectbox"] > div,
div[data-testid="stTextInput"] > div,
div[data-testid="stDateInput"] > div,
div[data-testid="stTimeInput"] > div,
div[data-testid="stNumberInput"] > div,
div[data-testid="stTextArea"] textarea {{
    border-radius: 0.45rem !important;
}}

/* Tabs */
.stTabs [data-baseweb="tab-list"] {{
    gap: 0.4rem;
    border-bottom: 1px solid {NSIA['border']};
}}
.stTabs [data-baseweb="tab"] {{
    color: {NSIA['text_secondary']};
    font-weight: 600;
    padding: 0.5rem 0.9rem;
    border-radius: 0.5rem 0.5rem 0 0;
}}
.stTabs [aria-selected="true"] {{
    color: {NSIA['navy']};
    background: {NSIA['white']};
    border-bottom: 2px solid {NSIA['gold']};
}}

/* Alerts / expander */
.stAlert {{
    border-radius: 0.6rem;
    border-left: 3px solid {NSIA['gold']};
    background-color: rgba(13,18,34,0.75);
    color: {NSIA['silver']};
}}
div[data-testid="stExpander"] {{
    border: 1px solid rgba(172,163,139,0.25);
    border-radius: 0.65rem;
    background-color: rgba(13,18,34,0.65);
    color: {NSIA['silver']};
}}

/* Separator */
hr {{
    border: none;
    border-top: 1px solid {NSIA['border']};
    margin: 1rem 0;
}}

/* Responsive */
@media (max-width: 768px) {{
    h1 {{ font-size: 1.4rem; }}
    h2 {{ font-size: 1.15rem; }}
    div[data-testid="stMetric"] {{
        padding: 0.7rem;
    }}
}}
</style>
""",
    unsafe_allow_html=True,
)

BASE_DIR = os.path.dirname(__file__)
EXCEL_FILE = os.path.join(BASE_DIR, "BDD_CALL_CENTER.xlsm")
DATA_DIR = os.path.join(BASE_DIR, os.pardir, "Fichiers Bruts")
CA_CUMUL_PATTERN = "CA CUMULE"

# ============================================================================
# DONNÉES NSIA
# ============================================================================

CALL_COLUMNS_NSIA = [
    "Date ",
    "TO",
    "Nom du Client",
    "numero de telephone",
    "Immatriculation",
    "Police",
    "Campagne ",
    "Reception",
    "Prise d'appel",
    "Produit existant",
    "Produit proposé",
    "Feedback",
    "CA",
    "Point de vente",
]

CALL_COLUMNS_EXTRA = [
    "Heure_appel",
    "Statut",
    "Motif_non_reponse",
    "Commentaire",
    "Satisfaction",
    "Recommendation",
    "Produit souhaite",
]

CALL_COLUMNS = CALL_COLUMNS_NSIA + CALL_COLUMNS_EXTRA

CALL_SHEET = "Call"
REF_SHEET = "References"

DEFAULT_CAMPAIGNS = ["SATURATION", "RELANCE", "RECUPERATION", "ONBOARDING", "RECEPTION"]
DEFAULT_PRISE_APPEL = ["Oui", "Non", "Injoignable"]
DEFAULT_PRODUCTS = [
    "AUTRES DOMMAGES AUX BIENS",
    "AUTOMOBILE",
    "BRIS DE MACHINES",
    "COMPLEMENTAIRE ACCIDENT DE TRAVAIL",
    "DOMMAGES CORPORELS",
    "FACULTES",
    "GLOBALE BANQUE",
    "GLOBALE DOMMAGES",
    "INDIVIDUELLE ACCIDENT GROUPE",
    "INDIVIDUELLE ACCIDENTS (GRPE RD)",
    "INDIVIDUELLE ACCIDENTS PARTICULIER",
    "INDIVIDUELLE VOYAGE",
    "INCENDIE RISQUES COMMERCIAUX",
    "INCENDIE RISQUES SIMPLES",
    "INDIVIDUELLE ACCIDENTS",
    "MULTIRISQUE PROFESSIONNELLE",
    "MULTIRISQUE HABITATION",
    "RC BICYCLETTE",
    "RC ASSOCIATIONS SPORTIVES",
    "RC CHEF D ENTREPRISE",
    "RC COLONIES DE VACANCES",
    "RC CHEF DE FAMILLE",
    "RC DECENNALE",
    "RC EXPLOITATION",
    "RC ENTREPRISES BATIMENT ET GENIE CIVIL",
    "RC ENTREPRISES INDUSTRIELLES ET COMMERCIALES",
    "RC EXPLOITATION BATIMENT ET GENIE CIVIL",
    "RC ORGANISATION DE MANIFESTATIONS",
    "RC PROFESSIONNELLE",
    "RC PROFESSIONS MEDICALES ET PARA-MEDICALE",
    "RC PROPRIETAIRE D IMMEUBLE",
    "RC SCOLAIRE",
    "RESPONSABILITE CIVILE",
    "SANTE",
    "TOUS RISQUES CHANTIERS",
    "TOUS RISQUES INFORMATIQUE",
    "TOUS RISQUES MONTAGES",
    "VOL",
]
DEFAULT_MOUVEMENTS = ["REN", "AFN"]
DEFAULT_RESEAUX = [
    "AGENTS GENERAUX",
    "AUTRES COURTIERS",
    "BANCASSURANCE",
    "BUREAUX DIRECTS",
    "GCA",
    "GRAS SAVOYE",
    "LA PROTECTRICE",
    "OLEA INSURANCE SOLUTIONS",
]
DEFAULT_RECEPTIONS = list(DEFAULT_MOUVEMENTS)
DEFAULT_TO = ["Audrey", "Sylvanus"]

DEFAULT_STATUS = [
    "Repondu",
    "Non repondu - Occupe",
    "Non repondu - Pas de reponse",
    "Non repondu - Messagerie",
    "Annule",
    "Termine",
]
DEFAULT_MOTIFS = [
    "Occupe",
    "Pas de reponse",
    "Messagerie vocal",
    "Numero invalide",
    "Client indisponible",
    "Appel transfere",
    "Autre",
]

DEFAULT_POINTS_DE_VENTE = list(DEFAULT_POINTS_DE_VENTE)

DEFAULT_FEEDBACK = [
    "Barrage",
    "Aucun numero",
    "Rappel apres",
    "Injoignable",
    "Difficulter de communication",
    "Refus",
    "Stop appel",
    "Sinistre en cours",
    "Faux numero",
]

# ============================================================================
# UTILITAIRES EXCEL
# ============================================================================


def _normalize_headers(headers):
    cleaned = []
    for h in headers:
        if h is None:
            cleaned.append(None)
        else:
            cleaned.append(str(h).strip())
    return cleaned


def lire_appels_excel():
    if not os.path.exists(EXCEL_FILE):
        return pd.DataFrame(columns=CALL_COLUMNS)

    try:
        wb = load_workbook(EXCEL_FILE, data_only=True)
        if CALL_SHEET not in wb.sheetnames:
            return pd.DataFrame(columns=CALL_COLUMNS)

        ws = wb[CALL_SHEET]
        raw_headers = [cell.value for cell in ws[1]]
        headers = _normalize_headers(raw_headers)

        data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None for v in row):
                data.append(row)

        df = pd.DataFrame(data, columns=headers if any(h is not None for h in headers) else CALL_COLUMNS)
        df.insert(0, "_excel_row", range(2, 2 + len(df)))

        rename_map = {
            "Date ": "Date",
            "Campagne ": "Campagne",
            "numero de telephone": "telephone",
        }
        df.rename(columns={k: v for k, v in rename_map.items() if k in df.columns}, inplace=True)

        if "Date" in df.columns:
            df["Date"] = pd.to_datetime(df["Date"], errors="coerce").dt.date
        if "CA" in df.columns:
            df["CA"] = pd.to_numeric(df["CA"], errors="coerce").fillna(0)

        return df

    except Exception as e:
        st.error(f"Erreur lors de la lecture du fichier Excel: {e}")
        return pd.DataFrame(columns=CALL_COLUMNS)


def sauvegarder_appel_excel(nouvel_appel):
    try:
        wb = load_workbook(EXCEL_FILE)

        if CALL_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(CALL_SHEET)
            ws.append(CALL_COLUMNS)
        else:
            ws = wb[CALL_SHEET]

        headers_existants = _normalize_headers([cell.value for cell in ws[1]])

        internal_to_nsia = {
            "Date": "Date ",
            "Campagne": "Campagne ",
            "telephone": "numero de telephone",
        }

        ligne = []
        for col in headers_existants:
            key = col
            if key in internal_to_nsia:
                key = internal_to_nsia[key]
            if key in nouvel_appel:
                ligne.append(nouvel_appel[key])
            else:
                ligne.append(None)

        headers_supp = []
        for c in CALL_COLUMNS_EXTRA:
            nsia = internal_to_nsia.get(c, c)
            if nsia not in headers_existants:
                headers_supp.append(nsia)

        if headers_supp:
            col_debut = ws.max_column + 1
            for i, col in enumerate(headers_supp):
                ws.cell(row=1, column=col_debut + i, value=col)
            for col in headers_supp:
                ligne.append(nouvel_appel.get(col, None))

        ws.append(ligne)
        wb.save(EXCEL_FILE)
        return True

    except Exception as e:
        st.error(f"Erreur lors de l'enregistrement: {e}")
        return False


def lire_references():
    refs = {
        "campagnes": list(DEFAULT_CAMPAIGNS),
        "to": list(DEFAULT_TO),
        "produits": list(DEFAULT_PRODUCTS),
        "receptions": list(DEFAULT_MOUVEMENTS),
        "prise_appel": list(DEFAULT_PRISE_APPEL),
        "reseau": list(DEFAULT_RESEAUX),
        "feedback": list(DEFAULT_FEEDBACK),
        "points_de_vente": list(DEFAULT_POINTS_DE_VENTE),
    }

    if not os.path.exists(EXCEL_FILE):
        return refs

    try:
        wb = load_workbook(EXCEL_FILE, data_only=True)
        if REF_SHEET not in wb.sheetnames:
            return refs

        ws = wb[REF_SHEET]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                v = str(row[0]).strip()
                if v and v not in refs["campagnes"]:
                    refs["campagnes"].append(v)
            if row[1]:
                v = str(row[1]).strip()
                if v and v not in refs["to"]:
                    refs["to"].append(v)
            if row[2]:
                v = str(row[2]).strip()
                if v and v not in refs["produits"]:
                    refs["produits"].append(v)
            if row[5]:
                v = str(row[5]).strip()
                if v and v not in refs["prise_appel"]:
                    refs["prise_appel"].append(v)
            if row[7]:
                v = str(row[7]).strip()
                if v and v not in refs["receptions"]:
                    refs["receptions"].append(v)
            if row[8]:
                v = str(row[8]).strip()
                if v and v not in refs["points_de_vente"]:
                    refs["points_de_vente"].append(v)

        return refs

    except Exception as e:
        st.warning(f"Impossible de lire les references: {e}")
        return refs


def ajouter_to_excel(nouveau_to):
    if not nouveau_to or not str(nouveau_to).strip():
        return False
    try:
        wb = load_workbook(EXCEL_FILE)
        if REF_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(REF_SHEET)
            ws.append(["Campagne", "TO", "Produit", None, None, "PriseAppel", None, "Reception"])
        else:
            ws = wb[REF_SHEET]

        existants = []
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
            if row[0]:
                existants.append(str(row[0]).strip())

        if str(nouveau_to).strip() in existants:
            return "existe"

        ws.append([None, str(nouveau_to).strip(), None, None, None, None, None, None])
        wb.save(EXCEL_FILE)
        return True
    except Exception as e:
        st.error(f"Erreur lors de l'ajout du TO: {e}")
        return False


def supprimer_to_excel(to_supprimer):
    if not to_supprimer:
        return False
    try:
        wb = load_workbook(EXCEL_FILE)
        if REF_SHEET not in wb.sheetnames:
            return False
        ws = wb[REF_SHEET]

        for row in ws.iter_rows(min_row=2):
            if row[1].value and str(row[1].value).strip() == str(to_supprimer).strip():
                row[1].value = None

        wb.save(EXCEL_FILE)
        return True
    except Exception as e:
        st.error(f"Erreur lors de la suppression du TO: {e}")
        return False


def ajouter_point_de_vente_excel(nouveau_pdv):
    if not nouveau_pdv or not str(nouveau_pdv).strip():
        return False
    try:
        wb = load_workbook(EXCEL_FILE)
        if REF_SHEET not in wb.sheetnames:
            ws = wb.create_sheet(REF_SHEET)
            ws.append(["Campagne", "TO", "Produit", None, None, "PriseAppel", None, "Reception", "PointDeVente"])
        else:
            ws = wb[REF_SHEET]

        existants = []
        for row in ws.iter_rows(min_row=2, min_col=9, max_col=9, values_only=True):
            if row[0]:
                existants.append(str(row[0]).strip())

        if str(nouveau_pdv).strip() in existants:
            return "existe"

        ws.append([None, None, None, None, None, None, None, None, str(nouveau_pdv).strip()])
        wb.save(EXCEL_FILE)
        return True
    except Exception as e:
        st.error(f"Erreur lors de l'ajout du point de vente: {e}")
        return False


def supprimer_point_de_vente_excel(pdv_supprimer):
    if not pdv_supprimer:
        return False
    try:
        wb = load_workbook(EXCEL_FILE)
        if REF_SHEET not in wb.sheetnames:
            return False
        ws = wb[REF_SHEET]

        for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
            if row[0].value and str(row[0].value).strip() == str(pdv_supprimer).strip():
                row[0].value = None

        wb.save(EXCEL_FILE)
        return True
    except Exception as e:
        st.error(f"Erreur lors de la suppression du point de vente: {e}")
        return False


def supprimer_ligne_excel_par_index(ligne_excel_index):
    if ligne_excel_index is None:
        return False
    try:
        wb = load_workbook(EXCEL_FILE)
        if CALL_SHEET not in wb.sheetnames:
            return False
        ws = wb[CALL_SHEET]
        max_row = ws.max_row
        if ligne_excel_index < 2 or ligne_excel_index > max_row:
            return False
        ws.delete_rows(ligne_excel_index)
        wb.save(EXCEL_FILE)
        return True
    except Exception as e:
        st.error(f"Erreur lors de la suppression de la ligne: {e}")
        return False


# ============================================================================
# PAGE : SAISIE D'APPEL
# ============================================================================


def page_saisie(refs):
    st.markdown("## Saisie d'appel")
    st.caption("Formulaire adapte a la campagne selectionnee")

    if "campagne_sel" not in st.session_state:
        st.session_state["campagne_sel"] = None

    campagne_sel = st.session_state["campagne_sel"]

    # ------------------------------------------------------------------
    # ÉTAPE 1 : Sélection de la campagne (chips sans emoji)
    # ------------------------------------------------------------------
    if not campagne_sel:
        st.markdown("### Choisissez la campagne")
        st.markdown("<div style='margin: 1rem 0;'></div>", unsafe_allow_html=True)

        camps = [
            ("ONBOARDING", "ONBOARDING"),
            ("SATURATION", "SATURATION"),
            ("RELANCE", "RELANCE"),
            ("RECUPERATION", "RECUPERATION"),
            ("RECEPTION", "RECEPTION"),
        ]

        cols = st.columns(len(camps))
        for i, (label, value) in enumerate(camps):
            with cols[i]:
                if st.button(label, key=f"chip_{value}", use_container_width=True):
                    st.session_state["campagne_sel"] = value
                    st.rerun()

        return

    # ------------------------------------------------------------------
    # ÉTAPE 2 : Formulaire adapté
    # ------------------------------------------------------------------
    col1, col2 = st.columns([3, 1], gap="medium")

    with col1:
        with st.form("form_appel", clear_on_submit=False):
            # Bloc 1 : Informations générales
            st.markdown(f"#### Informations generales — <span style='color:{NSIA['gold']}'>{campagne_sel}</span>", unsafe_allow_html=True)
            c1, c2 = st.columns(2)
            with c1:
                date_appel = st.date_input("Date", value=date.today())
                heure_appel = st.time_input("Heure", value=datetime.now().time())
                to_appel = st.selectbox("TO (Teleconseiller) *", options=refs["to"])
                nom_client = st.text_input("Nom du Client", placeholder="Ex: Koffi Ama")
            with c2:
                telephone = st.text_input("Numero de telephone", placeholder="Ex: 90 00 00 00")
                reception = st.selectbox("Reception / Mouvt *", options=refs["receptions"])
                point_vente = st.selectbox("Point de vente", options=[""] + refs["points_de_vente"])

            # Bloc 2 : Détails contrat / produit
            st.markdown("#### Details contrat / produit")
            if campagne_sel in ["ONBOARDING", "SATURATION"]:
                c3, c4 = st.columns(2)
                with c3:
                    immatriculation = st.text_input("Immatriculation", placeholder="Ex: TG 1234 AB")
                    police = st.text_input("Police / N contrat", placeholder="Ex: POL-2024-001")
                    produit_existant = st.selectbox("Produit existant", options=[""] + refs["produits"])
                with c4:
                    produit_propose = st.selectbox("Produit propose", options=[""] + refs["produits"])
                    prise_appel = st.selectbox("Prise d'appel", options=refs["prise_appel"])
            elif campagne_sel == "RECEPTION":
                c3, c4 = st.columns(2)
                with c3:
                    immatriculation = st.text_input("Immatriculation", placeholder="Ex: TG 1234 AB")
                    police = st.text_input("Police / N contrat", placeholder="Ex: POL-2024-001")
                with c4:
                    produit_souhaite = st.selectbox("Produit souhaite", options=[""] + refs["produits"])
                    prise_appel = st.selectbox("Prise d'appel", options=refs["prise_appel"])
                produit_existant = ""
                produit_propose = ""
            else:
                c3, c4 = st.columns(2)
                with c3:
                    immatriculation = st.text_input("Immatriculation", placeholder="Ex: TG 1234 AB")
                    police = st.text_input("Police / N contrat", placeholder="Ex: POL-2024-001")
                with c4:
                    prise_appel = st.selectbox("Prise d'appel", options=refs["prise_appel"])
                produit_existant = ""
                produit_propose = ""
                produit_souhaite = ""

            # Bloc 3 : Résultat et suivi
            st.markdown("#### Resultat et suivi")
            c5, c6 = st.columns(2)
            with c5:
                if campagne_sel == "ONBOARDING":
                    satisfaction = st.selectbox("Satisfaction", options=["", "Non satisfait", "Satisfait", "Neutre"])
                    recommendation = st.selectbox("Recommendation", options=["", "Neutre"] + [str(i) for i in range(1, 11)])
                    feedback = st.selectbox("Feedback", options=[""] + refs["feedback"])
                elif campagne_sel == "SATURATION":
                    feedback = st.selectbox("Feedback", options=[""] + refs["feedback"])
                    satisfaction = ""
                    recommendation = ""
                elif campagne_sel == "RECEPTION":
                    feedback = ""
                    satisfaction = ""
                    recommendation = ""
                else:
                    feedback = st.selectbox("Feedback", options=[""] + refs["feedback"])
                    satisfaction = ""
                    recommendation = ""

            with c6:
                ca = st.number_input("Chiffre d'affaires (FCFA)", min_value=0, value=0, step=1000)
                commentaire = st.text_area("Commentaire", placeholder="Notes supplementaires...", height=100)

            statut = ""

            submitted = st.form_submit_button("Enregistrer l'appel", type="primary", use_container_width=True)

            if submitted:
                nom_client_safe = nom_client.strip() if isinstance(nom_client, str) else nom_client
                telephone_safe = telephone.strip() if isinstance(telephone, str) else telephone
                nom_affiche = nom_client_safe or "Appel sans nom"
                nouvel_appel = {
                    "Date ": date_appel,
                    "TO": to_appel,
                    "Nom du Client": nom_client_safe or None,
                    "numero de telephone": telephone_safe or None,
                    "Immatriculation": immatriculation or None,
                    "Police": police or None,
                    "Campagne ": campagne_sel,
                    "Reception": reception,
                    "Prise d'appel": prise_appel,
                    "Produit existant": produit_existant or None,
                    "Produit proposé": produit_propose or None,
                    "Produit souhaite": produit_souhaite if campagne_sel == "RECEPTION" else None,
                    "Feedback": feedback or None,
                    "CA": ca,
                    "Point de vente": point_vente or None,
                    "Heure_appel": heure_appel.strftime("%H:%M"),
                    "Statut": statut or None,
                    "Motif_non_reponse": None,
                    "Commentaire": commentaire or None,
                    "Satisfaction": satisfaction or None,
                    "Recommendation": recommendation or None,
                }

                if sauvegarder_appel_excel(nouvel_appel):
                    st.success(f"Appel enregistre pour {nom_affiche} !")
                    st.balloons()
                    st.session_state["campagne_sel"] = None
                    st.rerun()

    with col2:
        st.markdown("### Campagne active")
        st.info(campagne_sel)

        if st.button("Changer de campagne", type="secondary", use_container_width=True):
            st.session_state["campagne_sel"] = None
            st.rerun()

        st.markdown("---")
        st.markdown("#### Guide de saisie")
        st.markdown(
            f"""
<div style="background-color:rgba(32,44,84,0.35);padding:1rem;border-radius:0.75rem;border-left:3px solid {NSIA['gold']};color:{NSIA['silver']};">
<ul style="margin:0;padding-left:1.1rem;">
<li><strong>Champs recommandes :</strong> Nom, Telephone</li>
<li>Enregistrement possible <strong>sans nom ni telephone</strong></li>
<li>Le formulaire s'adapte a la campagne choisie</li>
</ul>
</div>
""",
            unsafe_allow_html=True,
        )


# ============================================================================
# PAGE : TABLEAU DE BORD
# ============================================================================


def calculer_kpis(df):
    if df.empty:
        return {
            "nb_appels": 0,
            "nb_repondus": 0,
            "nb_non_repondus": 0,
            "taux_reponse": 0.0,
            "ca_total": 0,
            "ca_moyen": 0.0,
            "top_to": "N/A",
            "top_campagne": "N/A",
            "top_produit": "N/A",
        }

    nb_total = len(df)
    nb_repondus = len(df[df["Prise d'appel"] == "Oui"]) if "Prise d'appel" in df.columns else 0
    nb_non_repondus = nb_total - nb_repondus
    taux_reponse = (nb_repondus / nb_total * 100) if nb_total > 0 else 0.0
    ca_total = df["CA"].sum() if "CA" in df.columns else 0
    ca_moyen = ca_total / nb_total if nb_total > 0 else 0.0

    top_to = "N/A"
    if "TO" in df.columns and ca_total > 0:
        try:
            top_to = df.groupby("TO")["CA"].sum().idxmax()
        except ValueError:
            pass

    top_campagne = "N/A"
    if "Campagne" in df.columns and not df["Campagne"].dropna().empty:
        try:
            top_campagne = df["Campagne"].value_counts().idxmax()
        except ValueError:
            pass

    top_produit = "N/A"
    if "Produit proposé" in df.columns and ca_total > 0:
        try:
            produits_ca = df.groupby("Produit proposé")["CA"].sum()
            if not produits_ca.empty:
                top_produit = produits_ca.idxmax()
        except ValueError:
            pass

    return {
        "nb_appels": nb_total,
        "nb_repondus": nb_repondus,
        "nb_non_repondus": nb_non_repondus,
        "taux_reponse": round(taux_reponse, 1),
        "ca_total": ca_total,
        "ca_moyen": round(ca_moyen, 0),
        "top_to": top_to,
        "top_campagne": top_campagne,
        "top_produit": top_produit,
    }


def page_tableau_bord(df):
    st.header("Tableau de bord")

    if df.empty:
        st.warning("Aucune donnee disponible. Commencez par saisir des appels.")
        return

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        filtre_date_debut = st.date_input("Date debut", value=None)
    with c2:
        filtre_date_fin = st.date_input("Date fin", value=None)
    with c3:
        filtre_to = st.multiselect("TO", options=sorted(df["TO"].dropna().unique()))
    with c4:
        filtre_campagne = st.multiselect("Campagne", options=sorted(df["Campagne"].dropna().unique()))

    with st.expander("Autres filtres", expanded=False):
        c5, c6 = st.columns(2)
        with c5:
            filtre_reception = st.multiselect("Reception", options=sorted(df["Reception"].dropna().unique()))
        with c6:
            filtre_produit = st.multiselect("Produit propose", options=sorted(df.get("Produit proposé", pd.Series()).dropna().unique()))

    satisfaction_vals = sorted(df.get("Satisfaction", pd.Series()).dropna().unique().tolist())
    if satisfaction_vals:
        filtre_satisfaction = st.multiselect("Satisfaction", options=satisfaction_vals)
    else:
        filtre_satisfaction = []

    df_filtre = df.copy()
    if filtre_date_debut:
        df_filtre = df_filtre[df_filtre["Date"] >= filtre_date_debut]
    if filtre_date_fin:
        df_filtre = df_filtre[df_filtre["Date"] <= filtre_date_fin]
    if filtre_to:
        df_filtre = df_filtre[df_filtre["TO"].isin(filtre_to)]
    if filtre_campagne:
        df_filtre = df_filtre[df_filtre["Campagne"].isin(filtre_campagne)]
    if filtre_produit and "Produit proposé" in df_filtre.columns:
        df_filtre = df_filtre[df_filtre["Produit proposé"].isin(filtre_produit)]
    if filtre_satisfaction and "Satisfaction" in df_filtre.columns:
        df_filtre = df_filtre[df_filtre["Satisfaction"].isin(filtre_satisfaction)]

    if df_filtre.empty:
        st.warning("Aucune donnee correspondant aux filtres selectionnes.")
        return

    kpis = calculer_kpis(df_filtre)

    st.markdown("### Indicateurs cles")
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Nb Appels", kpis["nb_appels"])
    c2.metric("Repondus", kpis["nb_repondus"])
    c3.metric("Non Repondus", kpis["nb_non_repondus"])
    c4.metric("Taux Reponse", f"{kpis['taux_reponse']}%")
    c5.metric("CA Total (FCFA)", f"{kpis['ca_total']:,.0f}".replace(",", " "))

    c6, c7, c8 = st.columns(3)
    c6.metric("CA Moyen / Appel (FCFA)", f"{kpis['ca_moyen']:,.0f}".replace(",", " "))
    c7.metric("Top TO", kpis["top_to"])
    c8.metric("Top Campagne", kpis["top_campagne"])

    st.markdown("---")

    tab1, tab2, tab3, tab4 = st.tabs(["Campagnes & Produits", "Performance TO", "CA & Appels", "Details"])

    with tab1:
        c1, c2 = st.columns(2)
        with c1:
            camp_counts = df_filtre["Campagne"].value_counts().reset_index()
            camp_counts.columns = ["Campagne", "Nb Appels"]
            fig_camp = px.bar(camp_counts, x="Campagne", y="Nb Appels", color="Nb Appels", title="Nombre d'appels par campagne", text="Nb Appels")
            fig_camp.update_layout(showlegend=False)
            st.plotly_chart(fig_camp, use_container_width=True)
        with c2:
            camp_ca = df_filtre.groupby("Campagne")["CA"].sum().reset_index()
            camp_ca.columns = ["Campagne", "CA (FCFA)"]
            fig_ca = px.bar(camp_ca, x="Campagne", y="CA (FCFA)", color="CA (FCFA)", title="Chiffre d'affaires par campagne (FCFA)", text="CA (FCFA)")
            fig_ca.update_layout(showlegend=False)
            st.plotly_chart(fig_ca, use_container_width=True)

        c3, c4 = st.columns(2)
        with c3:
            prod_counts = df_filtre.get("Produit proposé", pd.Series()).value_counts().reset_index()
            prod_counts.columns = ["Produit", "Nb Appels"]
            fig_prod = px.pie(prod_counts, values="Nb Appels", names="Produit", title="Repartition des appels par produit propose")
            st.plotly_chart(fig_prod, use_container_width=True)
        with c4:
            prod_ca = df_filtre.groupby("Produit proposé")["CA"].sum().reset_index()
            prod_ca.columns = ["Produit", "CA (FCFA)"]
            fig_prod_ca = px.bar(prod_ca, x="Produit", y="CA (FCFA)", color="CA (FCFA)", title="CA par produit propose (FCFA)", text="CA (FCFA)")
            fig_prod_ca.update_layout(showlegend=False, xaxis_tickangle=-45)
            st.plotly_chart(fig_prod_ca, use_container_width=True)

    with tab2:
        c1, c2 = st.columns(2)
        with c1:
            to_counts = df_filtre["TO"].value_counts().reset_index()
            to_counts.columns = ["TO", "Nb Appels"]
            fig_to_count = px.bar(to_counts, x="TO", y="Nb Appels", color="Nb Appels", title="Nb appels par TO", text="Nb Appels")
            fig_to_count.update_layout(showlegend=False)
            st.plotly_chart(fig_to_count, use_container_width=True)
        with c2:
            to_ca = df_filtre.groupby("TO")["CA"].sum().reset_index()
            to_ca.columns = ["TO", "CA (FCFA)"]
            fig_to_ca = px.bar(to_ca, x="TO", y="CA (FCFA)", color="CA (FCFA)", title="CA par TO (FCFA)", text="CA (FCFA)")
            fig_to_ca.update_layout(showlegend=False)
            st.plotly_chart(fig_to_ca, use_container_width=True)

        to_stats = (
            df_filtre.groupby("TO")["Prise d'appel"]
            .apply(lambda x: (x == "Oui").sum() / len(x) * 100 if len(x) > 0 else 0)
            .reset_index()
        )
        to_stats.columns = ["TO", "Taux Reponse (%)"]
        fig_to_taux = px.bar(to_stats, x="TO", y="Taux Reponse (%)", color="Taux Reponse (%)", title="Taux de reponse par TO (%)", text="Taux Reponse (%)")
        fig_to_taux.update_layout(showlegend=False, yaxis_range=[0, 100])
        st.plotly_chart(fig_to_taux, use_container_width=True)

    with tab3:
        c1, c2 = st.columns(2)
        with c1:
            if "Date" in df_filtre.columns:
                df_daily = df_filtre.groupby("Date")["CA"].sum().reset_index().sort_values("Date")
                df_daily["CA Cumule"] = df_daily["CA"].cumsum()
                fig_ca_cum = px.line(df_daily, x="Date", y="CA Cumule", title="CA cumule dans le temps (FCFA)", markers=True)
                st.plotly_chart(fig_ca_cum, use_container_width=True)
        with c2:
            if "Date" in df_filtre.columns:
                daily_counts = df_filtre.groupby("Date").size().reset_index(name="Nb Appels").sort_values("Date")
                fig_daily = px.bar(daily_counts, x="Date", y="Nb Appels", title="Volume d'appels par jour", text="Nb Appels")
                fig_daily.update_layout(showlegend=False)
                st.plotly_chart(fig_daily, use_container_width=True)

        rec_stats = df_filtre.groupby("Reception").agg(Nb_Appels=("Reception", "size"), CA=("CA", "sum")).reset_index()
        rec_stats.columns = ["Reception", "Nb Appels", "CA (FCFA)"]
        c3, c4 = st.columns(2)
        with c3:
            fig_rec_count = px.pie(rec_stats, values="Nb Appels", names="Reception", title="Repartition REN vs AFN (nb appels)")
            st.plotly_chart(fig_rec_count, use_container_width=True)
        with c4:
            fig_rec_ca = px.pie(rec_stats, values="CA (FCFA)", names="Reception", title="Repartition REN vs AFN (CA FCFA)")
            st.plotly_chart(fig_rec_ca, use_container_width=True)

    with tab4:
        st.subheader("Statut des appels")
        status_counts = df_filtre["Statut"].value_counts().reset_index()
        status_counts.columns = ["Statut", "Nombre"]
        fig_status = px.bar(status_counts, x="Statut", y="Nombre", color="Nombre", title="Repartition par statut d'appel", text="Nombre")
        fig_status.update_layout(showlegend=False, xaxis_tickangle=-30)
        st.plotly_chart(fig_status, use_container_width=True)

        st.subheader("Feedbacks clients")
        fb_counts = df_filtre.get("Feedback", pd.Series()).dropna().value_counts().reset_index()
        fb_counts.columns = ["Feedback", "Nombre"]
        if not fb_counts.empty:
            fig_fb = px.bar(fb_counts, x="Feedback", y="Nombre", color="Nombre", title="Repartition des feedbacks", text="Nombre")
            fig_fb.update_layout(showlegend=False, xaxis_tickangle=-30)
            st.plotly_chart(fig_fb, use_container_width=True)
        else:
            st.info("Aucun feedback enregistre.")


# ============================================================================
# PAGE : HISTORIQUE
# ============================================================================


def page_historique(df):
    st.header("Historique des appels")

    if df.empty:
        st.warning("Aucune donnee disponible.")
        return

    c1, c2 = st.columns(2)
    with c1:
        filtre_date_debut = st.date_input("Date debut", value=None, key="hist_date_debut")
    with c2:
        filtre_date_fin = st.date_input("Date fin", value=None, key="hist_date_fin")

    recherche = st.text_input("Rechercher (nom client, telephone, police, immatriculation)", placeholder="Tapez un mot-cle...")

    satisfaction_vals_hist = sorted(df.get("Satisfaction", pd.Series()).dropna().unique().tolist())
    filtre_satisfaction_hist = st.multiselect("Satisfaction", options=satisfaction_vals_hist, key="filtre_satisfaction_hist") if satisfaction_vals_hist else []

    df_recherche = df.copy()
    if filtre_date_debut:
        df_recherche = df_recherche[df_recherche["Date"] >= filtre_date_debut]
    if filtre_date_fin:
        df_recherche = df_recherche[df_recherche["Date"] <= filtre_date_fin]
    if recherche:
        rech = recherche.lower()
        df_recherche = df_recherche[
            df_recherche["Nom du Client"].str.lower().str.contains(rech, na=False)
            | df_recherche["numero de telephone"].str.lower().str.contains(rech, na=False)
            | df_recherche.get("Police", pd.Series()).str.lower().str.contains(rech, na=False)
            | df_recherche.get("Immatriculation", pd.Series()).str.lower().str.contains(rech, na=False)
        ]

    if filtre_satisfaction_hist and "Satisfaction" in df_recherche.columns:
        df_recherche = df_recherche[df_recherche["Satisfaction"].isin(filtre_satisfaction_hist)]

    st.write(f"**{len(df_recherche)} appel(s) trouve(s)**")

    cols_affichage = [
        "Date",
        "Heure_appel",
        "TO",
        "Nom du Client",
        "numero de telephone",
        "Campagne",
        "Reception",
        "Prise d'appel",
        "Produit existant",
        "Produit proposé",
        "CA",
        "Statut",
        "Feedback",
        "Satisfaction",
        "Recommendation",
        "Point de vente",
    ]
    cols_dispo = [c for c in cols_affichage if c in df_recherche.columns]
    df_aff = df_recherche[cols_dispo].copy()

    if "Date" in df_aff.columns and "Heure_appel" in df_aff.columns:
        df_aff = df_aff.sort_values(by=["Date", "Heure_appel"], ascending=[False, False])

    selection = st.dataframe(
        df_aff,
        use_container_width=True,
        height=500,
        selection_mode="multi-row",
        on_select="rerun",
        key="hist_selection",
    )
    selected_rows = selection.get("selection", {}).get("rows", []) or []
    if selected_rows:
        st.caption(f"Ligne(s) selectionnee(s): {len(selected_rows)}")
        c1, c2 = st.columns(2)
        with c1:
            if st.button("Supprimer la selection", type="secondary", use_container_width=True, key="supprimer_historique"):
                excel_indices = [int(df_recherche.iloc[i]["_excel_row"]) for i in selected_rows if i < len(df_recherche)]
                success_count = 0
                for excel_index in sorted(set(excel_indices), reverse=True):
                    if supprimer_ligne_excel_par_index(excel_index):
                        success_count += 1
                if success_count:
                    st.success(f"{success_count} ligne(s) supprimee(s).")
                    st.rerun()
                else:
                    st.error("Aucune suppression possible.")
        with c2:
            if st.button("Rafraichir", type="primary", use_container_width=True, key="refresh_historique"):
                st.rerun()

    c1, c2 = st.columns(2)
    with c1:
        csv = df_aff.to_csv(index=False).encode("utf-8")
        st.download_button("Exporter en CSV", data=csv, file_name=f"appels_export_{date.today().isoformat()}.csv", mime="text/csv")
    with c2:
        import io
        output = io.BytesIO()
        workbook = pd.ExcelWriter(output, engine="xlsxwriter")
        df_aff.to_excel(workbook, sheet_name="Call", index=False)
        worksheet = workbook.sheets["Call"]
        worksheet.autofilter(0, 0, df_aff.shape[0], df_aff.shape[1] - 1)
        worksheet.freeze_panes(1, 0)
        header_fmt = workbook.book.add_format(
            {"bold": True, "bg_color": "#D9E1F2", "border": 1, "align": "center", "valign": "vcenter"}
        )
        for col_num, value in enumerate(df_aff.columns.values):
            worksheet.write(0, col_num, value, header_fmt)
            worksheet.set_column(col_num, col_num, max(12, min(40, len(str(value)) + 2)))
        workbook.close()
        st.download_button("Exporter en Excel", data=output.getvalue(), file_name=f"appels_export_{date.today().isoformat()}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ============================================================================
# PAGE : RECHERCHE CLIENT
# ============================================================================


def page_recherche_client(df):
    st.header("Recherche client")

    if df.empty:
        st.warning("Aucune donnee disponible.")
        return

    recherche = st.text_input("Rechercher par nom, police ou immatriculation", placeholder="Ex: Koffi / TG2000... / POL-2024...")

    df_recherche = df.copy()
    if recherche:
        rech = recherche.lower()
        df_recherche = df_recherche[
            df_recherche["Nom du Client"].str.lower().str.contains(rech, na=False)
            | df_recherche.get("Police", pd.Series()).str.lower().str.contains(rech, na=False)
            | df_recherche.get("Immatriculation", pd.Series()).str.lower().str.contains(rech, na=False)
        ]

    st.write(f"**{len(df_recherche)} resultat(s)**")

    if df_recherche.empty:
        st.info("Aucun client trouve.")
        return

    cols_affichage = [
        "Date",
        "Heure_appel",
        "TO",
        "Nom du Client",
        "numero de telephone",
        "Police",
        "Immatriculation",
        "Campagne",
        "Reception",
        "Prise d'appel",
        "Produit existant",
        "Produit proposé",
        "CA",
        "Statut",
        "Feedback",
        "Point de vente",
    ]
    cols_dispo = [c for c in cols_affichage if c in df_recherche.columns]
    df_aff = df_recherche[cols_dispo].copy()

    if "Date" in df_aff.columns and "Heure_appel" in df_aff.columns:
        df_aff = df_aff.sort_values(by=["Date", "Heure_appel"], ascending=[False, False])

    selection = st.dataframe(
        df_aff,
        use_container_width=True,
        height=500,
        selection_mode="multi-row",
        on_select="rerun",
        key="recherche_selection",
    )
    selected_rows = selection.get("selection", {}).get("rows", []) or []
    if selected_rows:
        st.caption(f"Ligne(s) selectionnee(s): {len(selected_rows)}")
        c1, c2 = st.columns(2)
        with c1:
            if st.button("Supprimer la selection", type="secondary", use_container_width=True, key="supprimer_recherche"):
                excel_indices = [int(df_recherche.iloc[i]["_excel_row"]) for i in selected_rows if i < len(df_recherche)]
                success_count = 0
                for excel_index in sorted(set(excel_indices), reverse=True):
                    if supprimer_ligne_excel_par_index(excel_index):
                        success_count += 1
                if success_count:
                    st.success(f"{success_count} ligne(s) supprimee(s).")
                    st.rerun()
                else:
                    st.error("Aucune suppression possible.")
        with c2:
            if st.button("Rafraichir", type="primary", use_container_width=True, key="refresh_recherche"):
                st.rerun()

    c1, c2 = st.columns(2)
    with c1:
        csv = df_aff.to_csv(index=False).encode("utf-8")
        st.download_button("Exporter en CSV", data=csv, file_name=f"recherche_client_{date.today().isoformat()}.csv", mime="text/csv")
    with c2:
        import io
        output = io.BytesIO()
        workbook = pd.ExcelWriter(output, engine="xlsxwriter")
        df_aff.to_excel(workbook, sheet_name="Recherche", index=False)
        worksheet = workbook.sheets["Recherche"]
        worksheet.autofilter(0, 0, df_aff.shape[0], df_aff.shape[1] - 1)
        worksheet.freeze_panes(1, 0)
        header_fmt = workbook.book.add_format(
            {"bold": True, "bg_color": "#D9E1F2", "border": 1, "align": "center", "valign": "vcenter"}
        )
        for col_num, value in enumerate(df_aff.columns.values):
            worksheet.write(0, col_num, value, header_fmt)
            worksheet.set_column(col_num, col_num, max(12, min(40, len(str(value)) + 2)))
        workbook.close()
        st.download_button("Exporter en Excel", data=output.getvalue(), file_name=f"recherche_client_{date.today().isoformat()}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ============================================================================
# PAGE : REFERENCES
# ============================================================================


def page_references(refs):
    st.header("Gestion des references")

    st.info("Visualisez et modifiez les listes de reference utilisees dans les formulaires.")

    c1, c2, c3 = st.columns(3)
    with c1:
        st.subheader("Campagnes")
        for c in refs["campagnes"]:
            st.write(f"- {c}")
        st.subheader("TOs")
        for t in refs["to"]:
            st.write(f"- {t}")
    with c2:
        st.subheader("Receptions")
        for r in refs["receptions"]:
            st.write(f"- {r}")
        st.subheader("Prise d'appel")
        for p in refs["prise_appel"]:
            st.write(f"- {p}")
    with c3:
        st.subheader("Produits")
        for p in refs["produits"]:
            st.write(f"- {p}")
        st.subheader("Statuts")
        for s in DEFAULT_STATUS:
            st.write(f"- {s}")


# ============================================================================
# PAGE : OPERATEURS
# ============================================================================


def page_operateurs(refs):
    st.header("Gestion des operateurs / TOs")

    st.markdown("Ajoutez ou supprimez des teleconseillers dans la base NSIA. Ces modifications sont écrites directement dans la feuille `References` du fichier Excel.")

    st.subheader("Operateurs existants")
    to_list = refs["to"]

    if not to_list:
        st.warning("Aucun operateur enregistre.")
    else:
        for to in to_list:
            c1, c2 = st.columns([3, 1])
            with c1:
                st.write(f"**{to}**")
            with c2:
                if st.button("Supprimer", key=f"del_{to}"):
                    result = supprimer_to_excel(to)
                    if result:
                        st.success(f"Operateur {to} supprime.")
                        st.rerun()
                    else:
                        st.error("Erreur lors de la suppression.")

    st.markdown("---")
    st.subheader("Ajouter un operateur")

    with st.form("form_ajout_to"):
        nouveau_to = st.text_input("Nom du nouvel operateur / TO", placeholder="Ex: Jean, Marie, Kofi...")
        submitted = st.form_submit_button("Ajouter l'operateur", type="primary", use_container_width=True)

        if submitted:
            if not nouveau_to or not str(nouveau_to).strip():
                st.error("Veuillez saisir un nom d'operateur.")
            else:
                result = ajouter_to_excel(nouveau_to)
                if result == True:
                    st.success(f"Operateur '{nouveau_to}' ajoute avec succes !")
                    st.rerun()
                elif result == "existe":
                    st.warning(f"L'operateur '{nouveau_to}' existe deja dans la base.")
                else:
                    st.error("Erreur lors de l'ajout.")


# ============================================================================
# PAGE : STATISTIQUES AVANCÉES
# ============================================================================


def page_statistiques(df):
    st.header("Statistiques avancees")

    if df.empty:
        st.warning("Aucune donnee disponible.")
        return

    c1, c2 = st.columns(2)
    with c1:
        filtre_date_debut = st.date_input("Date debut", value=None, key="stats_date_debut")
    with c2:
        filtre_date_fin = st.date_input("Date fin", value=None, key="stats_date_fin")

    df_stats = df.copy()
    if filtre_date_debut:
        df_stats = df_stats[df_stats["Date"] >= filtre_date_debut]
    if filtre_date_fin:
        df_stats = df_stats[df_stats["Date"] <= filtre_date_fin]

    if df_stats.empty:
        st.warning("Aucune donnee pour la periode selectionnee.")
        return

    kpis = calculer_kpis(df_stats)

    st.markdown("### KPIs globaux")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Nb Appels", kpis["nb_appels"])
    c2.metric("Taux Reponse (%)", f"{kpis['taux_reponse']}%")
    c3.metric("CA Total (FCFA)", f"{kpis['ca_total']:,.0f}".replace(",", " "))
    c4.metric("CA Moyen / Appel", f"{kpis['ca_moyen']:,.0f}".replace(",", " "))

    st.markdown("---")
    st.subheader("Performance par TO et par campagne")
    perf = (
        df_stats.groupby(["TO", "Campagne"])
        .agg(Appels=("TO", "size"), Repondus=("Prise d'appel", lambda x: (x == "Oui").sum()), CA=("CA", "sum"))
        .reset_index()
    )
    perf["Taux Reponse (%)"] = (perf["Repondus"] / perf["Appels"] * 100).round(1)
    perf.columns = ["TO", "Campagne", "Nb Appels", "Nb Repondus", "CA (FCFA)", "Taux Reponse (%)"]
    st.dataframe(perf, use_container_width=True, height=400)

    st.markdown("---")
    st.subheader("Export")
    csv = df_stats.to_csv(index=False).encode("utf-8")
    st.download_button("Exporter les donnees filtreees en CSV", data=csv, file_name=f"stats_avancees_{date.today().isoformat()}.csv", mime="text/csv")


# ============================================================================
# PAGE : POINTS DE VENTE
# ============================================================================


def page_points_de_vente(refs):
    st.header("Gestion des points de vente")

    st.markdown("Ajoutez ou supprimez des points de vente dans la base NSIA. Ces modifications sont écrites directement dans la feuille `References` du fichier Excel.")

    st.subheader("Points de vente existants")
    pdv_list = refs["points_de_vente"]

    if not pdv_list:
        st.warning("Aucun point de vente enregistre.")
    else:
        for pdv in pdv_list:
            c1, c2 = st.columns([3, 1])
            with c1:
                st.write(f"**{pdv}**")
            with c2:
                if st.button("Supprimer", key=f"del_pdv_{pdv}"):
                    result = supprimer_point_de_vente_excel(pdv)
                    if result:
                        st.success(f"Point de vente '{pdv}' supprime.")
                        st.rerun()
                    else:
                        st.error("Erreur lors de la suppression.")

    st.markdown("---")
    st.subheader("Ajouter un point de vente")

    with st.form("form_ajout_pdv"):
        nouveau_pdv = st.text_input("Nom du nouveau point de vente", placeholder="Ex: Lomé Bè, Agoè, Aného...")
        submitted = st.form_submit_button("Ajouter le point de vente", type="primary", use_container_width=True)

        if submitted:
            if not nouveau_pdv or not str(nouveau_pdv).strip():
                st.error("Veuillez saisir un nom de point de vente.")
            else:
                result = ajouter_point_de_vente_excel(nouveau_pdv)
                if result is True:
                    st.success(f"Point de vente '{nouveau_pdv}' ajoute avec succes !")
                    st.rerun()
                elif result == "existe":
                    st.warning(f"Le point de vente '{nouveau_pdv}' existe deja dans la base.")
                else:
                    st.error("Erreur lors de l'ajout.")


# ============================================================================
# PAGE : PARAMÈTRES
# ============================================================================


def page_parametres():
    st.header("Parametres")

    st.markdown("### Fichier Excel")
    st.code(EXCEL_FILE, language="text")

    st.markdown("### Dossiers")
    st.code(DATA_DIR, language="text")

    st.markdown("### Palette NSIA")
    for k, v in NSIA.items():
        st.markdown(f"- **{k}**: `{v}`")

    st.markdown("### Informations")
    st.markdown(f"- Page: **NSIA Call Center**")
    st.markdown(f"- Layout: **wide**")


# ============================================================================
# NAVIGATION PRINCIPALE
# ============================================================================


def main():
    st.set_page_config(page_title="NSIA Call Center", page_icon="📞", layout="wide", initial_sidebar_state="expanded")

    st.sidebar.title("NSIA Call Center")
    st.sidebar.markdown("**Assurances Togo**")
    st.sidebar.markdown("---")

    page = st.sidebar.radio(
        "Navigation",
        [
            "Saisie d'appel",
            "Tableau de bord",
            "Historique",
            "Recherche client",
            "Statistiques avancees",
            "Operateurs",
            "Points de vente",
            "References",
            "Parametres",
        ],
        label_visibility="collapsed",
    )

    st.sidebar.markdown("---")
    st.sidebar.caption(f"Fichier: `BDD_CALL_CENTER.xlsm`\nDerniere mise a jour: {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    df = lire_appels_excel()
    refs = lire_references()

    if page == "Saisie d'appel":
        page_saisie(refs)
    elif page == "Tableau de bord":
        page_tableau_bord(df)
    elif page == "Historique":
        page_historique(df)
    elif page == "Recherche client":
        page_recherche_client(df)
    elif page == "Statistiques avancees":
        page_statistiques(df)
    elif page == "Operateurs":
        page_operateurs(refs)
    elif page == "Points de vente":
        page_points_de_vente(refs)
    elif page == "References":
        page_references(refs)
    elif page == "Parametres":
        page_parametres()


if __name__ == "__main__":
    main()
